import json,io,sys,os,subprocess
from pathlib import Path
from hashlib import sha256
sys.path.append(str(Path.cwd()/'backend/.venv-office/lib/python3.12/site-packages'))
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from PIL import Image
base=Path('.local/batch3/b32b/evidence');dl=base/'browser-downloads';records=json.loads((base/'browser/all-network-before-restart.json').read_text());manifests={}
from workagent.db import Database
from workagent.service import Service,Principal
service=Service(Database());fixture=json.loads(Path('.local/batch3/b32b/fixture.json').read_text())
for r in records:
 if '/exports/' in r['url'] and r['url'].endswith('/download'):
  export_id=r['url'].split('/')[-2];m,_=service.download_export(Principal('local-human'),fixture['workspace_id'],export_id);manifests[m.format]=m.model_dump(mode='json')
results=[]
for fmt in ['docx','pdf','xlsx','png']:
 p=dl/('export.'+fmt);raw=p.read_bytes();m=manifests[fmt];assert sha256(raw).hexdigest()==m['blob']['sha256'];assert len(raw)==m['blob']['byte_length']
 result={'format':fmt,'path':str(p),'bytes':len(raw),'sha256':sha256(raw).hexdigest(),'manifest':m}
 if fmt=='docx':
  text='\n'.join(x.text for x in Document(io.BytesIO(raw)).paragraphs)
  assert 'Human web edit: use the revised figures.' in text and 'Human web annotation retained across restart.' in text and 'Exact quote: “Preserve this wording.”' in text
  result['text']=text
 if fmt=='pdf':
  reader=PdfReader(io.BytesIO(raw));text='\n'.join(x.extract_text() for x in reader.pages);assert 'Human web annotation retained across restart.' in text and 'Preserve this wording.' in text;result.update(pages=len(reader.pages),text=text)
 if fmt=='png':
  img=Image.open(io.BytesIO(raw));assert img.size==(1000,600);img.verify();assert [(x['category'],x['value']) for x in m['chart']['rows']]==[('East','0.3'),('West','1.5')];result['size']=[1000,600]
 if fmt=='xlsx':
  source=json.loads((base/'browser/workbook-source.json').read_text());book=load_workbook(io.BytesIO(raw),data_only=False);cache=load_workbook(io.BytesIO(raw),data_only=True);cells=[]
  for sheet in source['sheets']:
   for cell in sheet['cells']:
    got=book[sheet['name']][cell['address']];expected=cell['value'];kind=cell['kind']
    if kind=='formula':assert got.data_type=='f' and got.value==expected and cache[sheet['name']][cell['address']].value is None
    elif kind=='number':assert str(got.value)==expected or float(got.value)==float(expected)
    else:assert got.value==expected
    cells.append({'sheet':sheet['name'],'cell':cell['address'],'type':got.data_type,'value':got.value})
  assert any('Human xlsx imported copy annotation retained.' in str(c.value) for sh in book for row in sh for c in row)
  result['cells']=cells
 results.append(result)
(base/'download-readback.json').write_text(json.dumps(results,ensure_ascii=False,indent=2));print('Independently parsed and hash-bound',len(results),'browser downloads')
for fmt in ['docx','xlsx']:
 out=base/'rendered'/('browser-'+fmt);out.mkdir(exist_ok=True)
 subprocess.run(['libreoffice','-env:UserInstallation=file://'+str(Path.cwd()/'.local/batch3/b32b/lo-profile'),'--headless','--convert-to','pdf','--outdir',str(out),str(dl/('export.'+fmt))],check=True,timeout=60)
 subprocess.run(['pdftoppm','-scale-to','1100','-png','-singlefile',str(out/'export.pdf'),str(out/'page1')],check=True)
subprocess.run(['pdftoppm','-scale-to','1100','-png','-singlefile',str(dl/'export.pdf'),str(base/'rendered/browser-pdf')],check=True)
