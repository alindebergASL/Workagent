"""Standalone optional parser. ONLY execute inside document_sandbox's jail.

No service imports, credentials, host paths, URL resolution, or helper processes.
"""
import hashlib
import io
import json
import pathlib
import re
import stat
import sys
import zipfile

VERSION = 'confined-documents-v1'
LIMITS = {'input_bytes': 5242880, 'text_bytes': 102400, 'pages': 50,
          'sheets': 10, 'rows': 1000, 'cells': 10000, 'zip_bytes': 20971520,
          'zip_entries': 1000, 'zip_ratio': 100}

class Rejected(Exception):
    def __init__(self, reason, status='unsupported'):
        self.reason, self.status = reason, status

class Text:
    def __init__(self):
        self.parts, self.size, self.truncated = [], 0, False
    def add(self, value):
        raw = str(value).replace('\x00', '').encode('utf-8')
        room = LIMITS['text_bytes'] - self.size
        if len(raw) > room:
            self.truncated = True
            raw = raw[:room].decode('utf-8', 'ignore').encode('utf-8')
        self.parts.append(raw.decode('utf-8'))
        self.size += len(raw)
    def value(self):
        return ''.join(self.parts)


def preflight(raw, kind):
    from defusedxml import ElementTree as ET
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        infos = archive.infolist()
        if len(infos) > LIMITS['zip_entries'] or sum(i.file_size for i in infos) > LIMITS['zip_bytes']:
            raise Rejected('zip_expansion_limit')
        names = set()
        content_types = {}
        for info in infos:
            name = info.filename
            path = pathlib.PurePosixPath(name)
            if (name in names or name.startswith('/') or '\\' in name or ':' in name or '..' in path.parts
                    or stat.S_ISLNK(info.external_attr >> 16)):
                raise Rejected('unsafe_zip_path')
            names.add(name)
            if info.flag_bits & 1:
                raise Rejected('encrypted_zip', 'encrypted')
            if info.file_size > max(1, info.compress_size) * LIMITS['zip_ratio']:
                raise Rejected('zip_expansion_ratio')
            lower = name.lower()
            if (any(v in lower for v in ('vbaproject', 'oleobject', 'activex'))
                    or any(part in ('embeddings', 'externallinks') for part in path.as_posix().lower().split('/'))):
                raise Rejected('active_or_embedded_content')
        def xml(name):
            root = ET.fromstring(archive.read(name), forbid_dtd=True,
                                 forbid_entities=True, forbid_external=True)
            for node in root.iter():
                target = node.attrib.get('Target', '')
                if (node.attrib.get('TargetMode', '').lower() == 'external'
                        or (node.tag.endswith('}Relationship') and (target.startswith('//')
                            or re.match(r'^[A-Za-z][A-Za-z0-9+.-]*:', target)))):
                    raise Rejected('external_relationship')
                if any(v in str(node.attrib).lower() for v in ('macroenabled', 'vbaproject', 'oleobject')):
                    raise Rejected('active_or_embedded_content')
            return root

        ct_ns = '{http://schemas.openxmlformats.org/package/2006/content-types}'
        rel_ns = '{http://schemas.openxmlformats.org/package/2006/relationships}'
        office_rel = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument'
        if '[Content_Types].xml' not in names or '_rels/.rels' not in names:
            raise Rejected('format_container_mismatch')
        types = xml('[Content_Types].xml')
        if types.tag != ct_ns + 'Types':
            raise Rejected('format_container_mismatch')
        defaults, overrides = {}, {}
        for node in types:
            if node.tag == ct_ns + 'Default':
                key, mapping = node.attrib.get('Extension', '').lower(), defaults
            elif node.tag == ct_ns + 'Override':
                key, mapping = node.attrib.get('PartName', '').lower(), overrides
            else:
                raise Rejected('format_container_mismatch')
            if not key or key in mapping:
                raise Rejected('format_container_mismatch')
            mapping[key] = node.attrib.get('ContentType', '')
        for name in names:
            content_types[name] = overrides.get('/' + name.lower(), defaults.get(name.rsplit('.', 1)[-1].lower(), ''))
        # XML is determined by declared media type as well as conventional suffix.
        # Include unreferenced parts: a downstream parser may discover them later.
        relationships = None
        for name in names:
            media = content_types[name].split(';', 1)[0].strip().lower()
            if name.lower().endswith(('.xml', '.rels')) or media.endswith('+xml') or media in ('application/xml', 'text/xml'):
                root = xml(name)
                if name == '_rels/.rels':
                    relationships = root
        if relationships.tag != rel_ns + 'Relationships':
            raise Rejected('format_container_mismatch')
        ids = [node.attrib.get('Id') for node in relationships]
        if None in ids or len(set(ids)) != len(ids):
            raise Rejected('format_container_mismatch')
        main = [node for node in relationships if node.tag == rel_ns + 'Relationship'
                and node.attrib.get('Type') == office_rel]
        if len(main) != 1:
            raise Rejected('format_container_mismatch')
        target = main[0].attrib.get('Target', '')
        # Only canonical internal package names are supported. Never interpret
        # traversal, URI escapes, queries or fragments differently from the reader.
        part = target[1:] if target.startswith('/') else target
        if (not part or any(c in part for c in ('%', '?', '#', ':', '\\'))
                or any(p in ('', '.', '..') for p in part.split('/'))
                or main[0].attrib.get('TargetMode', 'Internal') != 'Internal'):
            raise Rejected('format_container_mismatch')
        expected = ('application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml' if kind == 'docx'
                    else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml')
        if part not in names or content_types.get(part) != expected:
            raise Rejected('format_container_mismatch')
        # Validate the actual relationship-selected main part, not a decoy at the
        # conventional path, even if its filename has no XML extension.
        root = xml(part)
        expected_root = ('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}document' if kind == 'docx'
                         else '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}workbook')
        if root.tag != expected_root:
            raise Rejected('format_container_mismatch')


def pdf(raw, text):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(raw), strict=True)
    if reader.is_encrypted:
        raise Rejected('encrypted_pdf', 'encrypted')
    total, covered, empty = len(reader.pages), 0, []
    for index in range(min(total, LIMITS['pages'])):
        value = reader.pages[index].extract_text() or ''
        if not value.strip():
            empty.append(index + 1)
        text.add(f'\n[Page {index + 1}]\n{value}\n')
        covered += 1
        if text.truncated:
            break
    coverage = {'total_pages': total, 'covered_pages': covered, 'omitted_pages': total - covered,
                'empty_pages_needing_ocr': empty, 'omissions': ['images, layout and non-text content'],
                'partial_last_page': text.truncated}
    return coverage, total > covered


def docx(raw, text):
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    document = Document(io.BytesIO(raw))
    paragraphs = tables = 0
    for child in document.element.body:
        if child.tag.endswith('}p'):
            paragraph = Paragraph(child, document)
            bullet = '[list] ' if paragraph._p.xpath('./w:pPr/w:numPr') or 'List' in (getattr(paragraph.style, 'name', '') or '') else ''
            text.add(bullet + paragraph.text + '\n')
            paragraphs += 1
        elif child.tag.endswith('}tbl'):
            table = Table(child, document)
            text.add('[Table]\n')
            for row in table.rows:
                text.add('\t'.join(cell.text for cell in row.cells) + '\n')
                if text.truncated:
                    break
            tables += 1
        if text.truncated:
            break
    total_paragraphs = len(document.element.body.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p'))
    total_tables = len(document.element.body.findall('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tbl'))
    return {'paragraphs_covered': paragraphs, 'tables_covered': tables,
            'total_paragraphs': total_paragraphs, 'omitted_paragraphs': total_paragraphs - paragraphs,
            'total_tables': total_tables, 'omitted_tables': total_tables - tables,
            'omissions': ['headers', 'footers', 'comments', 'tracked changes',
                          'nested tables, drawings and exact list numbering/layout'],
            'partial_last_block': text.truncated}, False


def xlsx(raw, text):
    from openpyxl import load_workbook
    # Normal mode exposes merged ranges; ZIP bounds + AS cap contain allocation.
    book = load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(raw), data_only=True, keep_links=False)
    sheets, cells, rows, clipped = [], 0, 0, False
    for sheet in book.worksheets[:LIMITS['sheets']]:
        info = {'name': sheet.title, 'state': sheet.sheet_state, 'total_rows': sheet.max_row,
                'total_columns': sheet.max_column, 'covered_rows': 0,
                'visited_cells': 0, 'last_cell': None, 'partial_last_row': False,
                'merged_ranges': sorted(str(r) for r in sheet.merged_cells.ranges)[:100],
                'omitted_merged_ranges': max(0, len(sheet.merged_cells.ranges) - 100)}
        text.add(f'\n[Sheet {sheet.title}; {sheet.sheet_state}]\n')
        for row in sheet.iter_rows(max_row=min(sheet.max_row, LIMITS['rows'] - rows),
                                   max_col=min(sheet.max_column, LIMITS['cells'])):
            if rows >= LIMITS['rows'] or cells >= LIMITS['cells'] or text.truncated:
                clipped = True
                break
            last_column = 0
            for cell in row:
                if cells >= LIMITS['cells'] or text.truncated:
                    clipped = True
                    break
                cells += 1
                info['visited_cells'] += 1
                info['last_cell'] = cell.coordinate
                last_column = cell.column
                if cell.value is not None:
                    value = str(cell.value)
                    if cell.data_type == 'f':
                        value = f'inert formula {value}; cached={cached[sheet.title][cell.coordinate].value!s}; references not resolved'
                    text.add(f'{cell.coordinate}: {value}\n')
            info['partial_last_row'] = text.truncated or last_column < sheet.max_column
            rows += 1
            info['covered_rows'] += 1
        if info['covered_rows'] < sheet.max_row or sheet.max_column > LIMITS['cells']:
            clipped = True
        info['omitted_rows'] = sheet.max_row - info['covered_rows']
        info['omitted_cell_positions'] = sheet.max_row * sheet.max_column - info['visited_cells']
        sheets.append(info)
        if cells >= LIMITS['cells'] or rows >= LIMITS['rows'] or text.truncated:
            break
    omitted = len(book.worksheets) - len(sheets)
    return {'total_sheets': len(book.worksheets), 'covered_sheets': len(sheets),
            'omitted_sheets': omitted, 'sheets': sheets, 'covered_rows': rows,
            'visited_cells': cells, 'omissions': ['formatting, charts, comments; formula references unsupported; no recalculation']}, clipped or omitted > 0


def extract(raw, kind):
    text = Text()
    coverage = {'omissions': []}
    try:
        # Dependency check happens in the jail, never import optional parsers in service.
        import pypdf, docx as docx_package, openpyxl, defusedxml  # noqa: F401
        if kind != 'pdf':
            preflight(raw, kind)
        coverage, clipped = {'pdf': pdf, 'docx': docx, 'xlsx': xlsx}[kind](raw, text)
        status = 'truncated' if clipped or text.truncated else 'extracted'
        reason = 'bounded_extraction' if status == 'truncated' else 'text_extracted'
        if kind == 'pdf' and coverage['covered_pages'] == len(coverage['empty_pages_needing_ocr']):
            status, reason, text = 'needs_ocr', 'covered_pages_need_ocr_or_are_empty', Text()
    except ImportError:
        status, reason, text = 'unavailable', 'optional_dependencies_missing', Text()
    except Rejected as exc:
        status, reason, text = exc.status, exc.reason, Text()
    except Exception:
        status, reason, text = 'failed', 'malformed_or_resource_limited_document', Text()
    value = text.value()
    return {'parser_version': VERSION, 'status': status, 'reason': reason, 'text': value,
            'text_sha256': hashlib.sha256(value.encode()).hexdigest(),
            'original_sha256': hashlib.sha256(raw).hexdigest(), 'limits': LIMITS, 'coverage': coverage}

if __name__ == '__main__':
    # -I disables PYTHONPATH; mount only selected optional distributions at /packages.
    sys.path.insert(0, '/packages')
    raw = pathlib.Path('/input').read_bytes()
    if len(raw) > LIMITS['input_bytes']:
        raise SystemExit(2)
    print(json.dumps(extract(raw, sys.argv[1]), ensure_ascii=True, separators=(',', ':')))
