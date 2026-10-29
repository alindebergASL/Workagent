"""Fixed trusted conversion program. Production execution ONLY in office_sandbox.

All output uses bounded in-memory bytes and one JSON protocol record. No output
filenames/paths, arbitrary code, evaluation, shell, URL or plugin inputs exist.
"""
import base64
import csv
from decimal import Decimal, localcontext, ROUND_HALF_EVEN
import hashlib
import io
import json
from pathlib import Path
import sys

if not __package__:
    sys.path.insert(0, '/packages')

if __package__:
    from . import office_content as C
    from .document_parser import preflight
else:
    import office_content as C
    from document_parser import preflight

MIMES = {'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
         'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
         'pdf': 'application/pdf', 'png': 'image/png'}


def import_office(raw, mime):
    if mime == MIMES['pdf']:
        if not raw.startswith(b'%PDF-'): raise ValueError('PDF signature')
        if __package__:
            from .document_parser import pdf, Text
        else:
            from document_parser import pdf, Text
        extracted=Text()
        coverage,clipped=pdf(raw,extracted)
        if clipped or extracted.truncated or coverage['empty_pages_needing_ocr']:
            raise ValueError('PDF needs OCR or exceeds extraction limits')
        return C.validate(dict(office_version=C.VERSION,kind='document',paragraphs=extracted.value().splitlines(),annotations=[]))
    kind = next((k for k in ('docx', 'xlsx') if MIMES[k] == mime), None)
    if kind is None: raise ValueError('unsupported original')
    preflight(raw, kind)
    if kind == 'docx':
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(io.BytesIO(raw))
        paragraphs = []
        for child in doc.element.body:
            if child.tag.endswith('}p'):
                paragraphs.append(Paragraph(child, doc).text)
            elif child.tag.endswith('}tbl'):
                # Plain row values remain editable; no claim of original table fidelity.
                paragraphs.extend('\t'.join(c.text for c in row.cells) for row in Table(child, doc).rows)
            if len(paragraphs) > 2000: raise ValueError('paragraph limit')
        return C.validate(dict(office_version=C.VERSION, kind='document', paragraphs=paragraphs, annotations=[]))
    from openpyxl import load_workbook
    from datetime import date, datetime, time
    book = load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
    cache = load_workbook(io.BytesIO(raw), data_only=True, keep_links=False)
    if len(book.worksheets) > 10: raise ValueError('sheet limit')
    sheets, total = [], 0
    for sheet in book.worksheets:
        if sheet.max_row > 999 or sheet.max_column > 702: raise ValueError('sheet dimension limit')
        total += sheet.max_row * sheet.max_column
        if total > 10000: raise ValueError('cell limit')
        cells = []
        for row in sheet.iter_rows():
            for cell in row:
                v, cached = cell.value, None
                if v is None: continue
                if cell.data_type == 'f':
                    k, v = 'formula', C.formula(v)
                    cached = cache[sheet.title][cell.coordinate].value
                    if cached is not None and type(cached) is not bool: cached = str(cached)
                elif type(v) is bool: k = 'boolean'
                elif isinstance(v, (int, float, Decimal)):
                    k, v = 'number', str(v)
                elif isinstance(v, (date, datetime, time)):
                    k, v = 'string', v.isoformat()
                else: k, v = 'string', str(v)
                cells.append(dict(address=cell.coordinate, kind=k, value=v, cached_value=cached))
        sheets.append(dict(name=sheet.title, cells=cells))
    return C.validate(dict(office_version=C.VERSION, kind='workbook', sheets=sheets, annotations=[]))


def tabular(raw, mime):
    if mime == 'text/csv':
        rows = list(csv.reader(io.StringIO(raw.decode('utf-8-sig'), newline=''), strict=True))
        if not 2 <= len(rows) <= 1001 or not 1 <= len(rows[0]) <= 100:
            raise ValueError('table size')
        headers, values = rows[0], rows[1:]
    else:
        value = json.loads(raw)
        C.validate(value)
        if value['kind'] != 'workbook' or len(value['sheets']) != 1:
            raise ValueError('single typed table required')
        from openpyxl.utils.cell import coordinate_to_tuple
        cells = {coordinate_to_tuple(c['address']): c for c in value['sheets'][0]['cells']}
        height = max((r for r, _ in cells), default=0)
        width = max((c for _, c in cells), default=0)
        if not 2 <= height <= 1000 or not 1 <= width <= 100: raise ValueError('table size')
        def val(r, c):
            cell = cells.get((r,c))
            if cell is None or cell['kind'] == 'blank': return None
            if cell['kind'] in ('formula', 'boolean'): raise ValueError('numeric/string table only')
            return cell['value']
        headers = [val(1, c) for c in range(1, width+1)]
        values = [[val(r,c) for c in range(1,width+1)] for r in range(2,height+1)]
    if any(not isinstance(h,str) or not h or len(h)>100 for h in headers) or len(set(headers)) != len(headers):
        raise ValueError('unique named columns required')
    if any(len(row) != len(headers) for row in values): raise ValueError('ragged table')
    return headers, values


def aggregate(raw, mime, spec):
    keys = {'category_column','value_column','aggregation','title','x_label','y_label'}
    if not isinstance(spec,dict) or set(spec) != keys: raise ValueError('chart shape')
    for key in ('category_column','title','x_label','y_label'): C.text(spec[key], 200)
    op = spec['aggregation']
    if op not in ('sum','mean','count') or (op=='count') != (spec['value_column'] is None):
        raise ValueError('chart aggregation')
    headers, rows = tabular(raw,mime)
    cat = headers.index(spec['category_column'])
    val = headers.index(spec['value_column']) if op != 'count' else None
    groups = {}
    with localcontext() as ctx:
        ctx.prec = 160  # Versioned finite precision, disclosed in ChartResult.
        ctx.rounding = ROUND_HALF_EVEN
        for row in rows:
            category = row[cat]
            if not isinstance(category,str) or not category or len(category)>100: raise ValueError('missing category')
            n = Decimal(1) if val is None else C.decimal(row[val])
            total, count = groups.get(category,(Decimal(0),0))
            groups[category] = (total+n,count+1)
            if len(groups)>50: raise ValueError('group limit')
        return [dict(category=k,value=str(v/n if op=='mean' else Decimal(n) if op=='count' else v),count=n)
                for k,(v,n) in sorted(groups.items())]


def content(raw, mime):
    if mime == 'application/json':
        value = json.loads(raw)
        if isinstance(value,dict) and value.get('office_version'):
            return C.validate(value)
    if mime == 'text/csv':
        headers, rows = tabular(raw,mime)
        from openpyxl.utils.cell import get_column_letter
        cells = [dict(address=f'{get_column_letter(j)}{i}', kind='string', value=v, cached_value=None)
                 for i,row in enumerate([headers,*rows],1) for j,v in enumerate(row,1)]
        return C.validate(dict(office_version=C.VERSION,kind='workbook',sheets=[dict(name='Data',cells=cells)],annotations=[]))
    return C.validate(dict(office_version=C.VERSION,kind='document',paragraphs=raw.decode('utf-8-sig').splitlines(),annotations=[]))


def paragraphs(value):
    if value['kind']=='document': return list(value['paragraphs'])
    return [line for s in value['sheets'] for line in [s['name'], *[
        c['address']+': '+str(c['value'])+(' [formula; not computed; supplied/imported cache, unverified='+str(c['cached_value'])+']' if c['kind']=='formula' else '')
        for c in s['cells']]]]


def render(raw, mime, fmt, spec, omissions):
    out, chart = io.BytesIO(), None
    if fmt == 'png':
        rows = aggregate(raw,mime,spec)
        chart = dict(computation_version='decimal-grouped-v1',input_sha256=hashlib.sha256(raw).hexdigest(),
                     spec=spec,rows=rows,null_handling='reject',decimal_precision=160,rounding='ROUND_HALF_EVEN')
        import matplotlib
        matplotlib.use('Agg')
        from matplotlib.figure import Figure
        from matplotlib.backends.backend_agg import FigureCanvasAgg
        fig = Figure(figsize=(10,6),dpi=100)
        FigureCanvasAgg(fig)
        ax = fig.subplots()
        ax.bar([r['category'] for r in rows],[float(Decimal(r['value'])) for r in rows])
        ax.set(title=spec['title'],xlabel=spec['x_label'],ylabel=spec['y_label'])
        ax.tick_params(axis='x',rotation=45)
        fig.tight_layout()
        fig.savefig(out,format='png',metadata={'Software':'Workagent decimal-grouped-v1'})
        omissions.append('Chart pixels omit input annotations and source details; inspect this manifest and original input. Numeric rows are decimal results; plot coordinates use floating point.')
        return out.getvalue(),chart
    value = content(raw,mime)
    lines = paragraphs(value) + ['Human annotation: '+n for n in value['annotations']]
    lines += ['OMITTED: '+n for n in omissions]
    if fmt == 'docx':
        from docx import Document
        doc = Document()
        for line in lines: doc.add_paragraph(line)
        doc.save(out)
    elif fmt == 'pdf':
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from xml.sax.saxutils import escape
        # Fixed system font path mounted read-only with /usr, never user-selected.
        font = TTFont('OfficeSans','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf')
        pdfmetrics.registerFont(font)
        if any(ord(c) not in font.face.charWidths for line in lines for c in line if c not in '\n\r\t'):
            raise ValueError('unsupported PDF glyph')
        style = ParagraphStyle('body',fontName='OfficeSans',fontSize=10,leading=14,wordWrap='CJK')
        flow = []
        for line in lines:
            flow.extend([Paragraph(escape(line).replace('\n','<br/>').replace('\t','    '),style),Spacer(1,6)])
        SimpleDocTemplate(out, pagesize=(595,842),leftMargin=42,rightMargin=42,topMargin=42,bottomMargin=42).build(flow)
    elif fmt == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.workbook.properties import CalcProperties
        book = Workbook()
        book.remove(book.active)
        book.calculation = CalcProperties(calcMode='manual',fullCalcOnLoad=False,forceFullCalc=False)
        if value['kind']=='workbook':
            for sheet in value['sheets']:
                dest = book.create_sheet(sheet['name'])
                for cell in sheet['cells']:
                    v,k = cell['value'],cell['kind']
                    if k=='number': v=C.decimal(v)
                    target = dest[cell['address']]
                    target.value=v
                    if k=='string': target.data_type='s'  # '=...' stays string data
        else:
            dest = book.create_sheet('Document')
            for i,line in enumerate(value['paragraphs'],1):
                dest.cell(i,1,line).data_type='s'
        name = 'Export notes'
        while name in book.sheetnames: name += '_'
        notes = book.create_sheet(name)
        for i,line in enumerate(['Human annotation: '+n for n in value['annotations']]+['OMITTED: '+n for n in omissions],1):
            notes.cell(i,1,line).data_type='s'
        book.save(out)
    else: raise ValueError('unsupported format')
    return out.getvalue(),chart


def readback(raw,fmt):
    # Read actual emitted bytes before reporting success, inside the same jail.
    if fmt in ('docx','xlsx'):
        preflight(raw,fmt)
        if fmt=='docx':
            from docx import Document
            Document(io.BytesIO(raw))
        else:
            from openpyxl import load_workbook
            check=load_workbook(io.BytesIO(raw),data_only=False,keep_links=False)
            if not check.worksheets: raise ValueError('empty workbook')
            check.close()
    elif fmt=='pdf':
        from pypdf import PdfReader
        check=PdfReader(io.BytesIO(raw),strict=True)
        if not 1<=len(check.pages)<=50: raise ValueError('PDF page bound')
    elif fmt=='png':
        from PIL import Image
        image=Image.open(io.BytesIO(raw))
        image.verify()
        if image.size!=(1000,600): raise ValueError('chart dimension bound')


def convert(job):
    expected={'operation','input','input_sha256','mime_type','format','chart_spec','source_omissions'}
    if not isinstance(job,dict) or set(job)!=expected: raise ValueError('protocol')
    raw=base64.b64decode(job['input'],validate=True)
    if len(raw)>5242880 or hashlib.sha256(raw).hexdigest()!=job['input_sha256']: raise ValueError('input integrity')
    omissions=list(C.OMISSIONS)+job['source_omissions']
    if job['operation']=='import':
        value=import_office(raw,job['mime_type'])
        output=json.dumps(value,ensure_ascii=False,separators=(',',':')).encode()
        mime,chart='application/json',None
    elif job['operation']=='render':
        if len(raw)>200000: raise ValueError('input size')
        if (job['format']=='png') != (job['chart_spec'] is not None): raise ValueError('chart format')
        output,chart=render(raw,job['mime_type'],job['format'],job['chart_spec'],omissions)
        mime=MIMES[job['format']]
    else: raise ValueError('operation')
    if not 0<len(output)<=1048576: raise ValueError('output size')
    if job['operation']=='render': readback(output,job['format'])
    return dict(version='office-converter-v1',input_sha256=job['input_sha256'],
                output=base64.b64encode(output).decode(),sha256=hashlib.sha256(output).hexdigest(),
                mime_type=mime,omissions=omissions,chart=chart)


if __name__=='__main__':
    # -I/-S: no ambient site imports. Only reviewed mounted dependencies and code.
    try:
        result=convert(json.loads(Path('/input').read_bytes()))
        sys.stdout.write(json.dumps(result,ensure_ascii=True,separators=(',',':')))
    except Exception:
        # Never print exception messages, document content, library stderr or paths.
        raise SystemExit(2)
