"""Real independent DOCX/XLSX readback; confined boundary tested without claiming jail proof."""
import base64
from contextlib import contextmanager
from hashlib import sha256
import io
import json
from pathlib import Path
import zipfile
import pytest
from workagent import office_content as C, office_converter as converter, office_sandbox as sandbox
from workagent.errors import DomainError
from workagent.office_models import ChartSpec
from workagent.workspace_file_models import ExportFile


def job(raw,mime,fmt=None,operation='render',spec=None):
    return dict(operation=operation,input=base64.b64encode(raw).decode(),input_sha256=sha256(raw).hexdigest(),
                mime_type=mime,format=fmt,chart_spec=spec,source_omissions=[])


def workbook():
    return dict(office_version='office-v1',kind='workbook',sheets=[dict(name='Data',cells=[
        dict(address='A1',kind='string',value='=HYPERLINK("https://example.invalid")',cached_value=None),
        dict(address='B1',kind='number',value='2.50',cached_value=None),
        dict(address='B2',kind='number',value='-3',cached_value=None),
        dict(address='B3',kind='formula',value='=SUM(B1:B2)',cached_value='-0.5')])],annotations=['Human note — keep exact'])


def test_real_xlsx_explicit_formula_string_and_annotations():
    from openpyxl import load_workbook
    data=converter.convert(job(json.dumps(workbook()).encode(),'application/json','xlsx'))
    raw=base64.b64decode(data['output'])
    book=load_workbook(io.BytesIO(raw),data_only=False)
    assert book['Data']['A1'].data_type=='s'
    assert book['Data']['A1'].value.startswith('=HYPERLINK')
    assert book['Data']['B1'].value==2.5 and book['Data']['B2'].value==-3
    assert book['Data']['B3'].data_type=='f' and book['Data']['B3'].value=='=SUM(B1:B2)'
    assert 'Human note — keep exact' in book['Export notes']['A1'].value
    assert load_workbook(io.BytesIO(raw),data_only=True)['Data']['B3'].value is None
    assert any('not evaluated' in s for s in data['omissions'])
    assert sha256(raw).hexdigest()==data['sha256']


def test_real_docx_quotes_and_annotations_roundtrip():
    from docx import Document
    value=dict(office_version='office-v1',kind='document',paragraphs=['An attributed quotation: “Keep this exact.”','Second paragraph.'],annotations=['Protected human note'])
    data=converter.convert(job(json.dumps(value).encode(),'application/json','docx'))
    raw=base64.b64decode(data['output'])
    paragraphs=[p.text for p in Document(io.BytesIO(raw)).paragraphs]
    assert paragraphs[:2]==value['paragraphs']
    assert 'Human annotation: Protected human note' in paragraphs
    imported=converter.convert(job(raw,converter.MIMES['docx'],operation='import'))
    assert json.loads(base64.b64decode(imported['output']))['paragraphs'][:2]==value['paragraphs']


def test_import_original_formula_and_cached_values():
    from openpyxl import Workbook
    book=Workbook();sheet=book.active;sheet['A1']=4;sheet['A2']='=SUM(A1:A1)'
    out=io.BytesIO();book.save(out)
    # An external producer's historical cache, independently injected into OOXML.
    rewritten=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(out.getvalue())) as z,zipfile.ZipFile(rewritten,'w',zipfile.ZIP_DEFLATED) as dest:
        for name in z.namelist():
            raw=z.read(name)
            if name=='xl/worksheets/sheet1.xml': raw=raw.replace(b'<f>SUM(A1:A1)</f><v></v>',b'<f>SUM(A1:A1)</f><v>4</v>')
            dest.writestr(name,raw)
    imported=converter.import_office(rewritten.getvalue(),converter.MIMES['xlsx'])
    cell=imported['sheets'][0]['cells'][1]
    assert cell['kind']=='formula' and cell['value']=='=SUM(A1:A1)' and cell['cached_value']=='4'


@pytest.mark.parametrize('value',['=HYPERLINK("http://secret")','=WEBSERVICE(A1)',"='[external.xlsx]Sheet'!A1",'=cmd|A1!A2','=INDIRECT(A1)','=A1000','=SUM(Sheet1!A1)', '=SUM(A1)&"secret"'])
def test_unsafe_explicit_formula_rejected(value):
    with pytest.raises(ValueError): C.formula(value)


def test_decimal_aggregation_actual_rows_and_null_rejection():
    spec=dict(category_column='team',value_column='hours',aggregation='sum',title='Hours',x_label='Team',y_label='Hours')
    raw=b'team,hours\na,0.1\na,0.2\nb,-2.5\n'
    assert converter.aggregate(raw,'text/csv',spec)==[dict(category='a',value='0.3',count=2),dict(category='b',value='-2.5',count=1)]
    assert converter.aggregate(raw,'text/csv',{**spec,'aggregation':'mean'})[0]['value']=='0.15'
    assert converter.aggregate(raw,'text/csv',{**spec,'aggregation':'count','value_column':None})[0]['value']=='2'
    for bad in [b'team,hours\na,\n',b'team,hours\na,NaN\n',b'team,hours\na,=1+2\n',b'team,team\na,1\n']:
        with pytest.raises((ValueError,TypeError)): converter.aggregate(bad,'text/csv',spec)


def test_zip_bomb_external_relationship_and_wrong_mime():
    from openpyxl import Workbook
    stream=io.BytesIO();Workbook().save(stream)
    raw=stream.getvalue()
    with pytest.raises(Exception): converter.import_office(raw,converter.MIMES['docx'])
    for name,data in [('bomb.xml',b'0'*100000),('bad.rels',b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship TargetMode="External" Target="https://secret.invalid" /></Relationships>')]:
        stream=io.BytesIO(raw)
        with zipfile.ZipFile(stream,'a',zipfile.ZIP_DEFLATED) as z: z.writestr(name,data)
        with pytest.raises(Exception): converter.import_office(stream.getvalue(),converter.MIMES['xlsx'])


def test_wrong_input_hash_and_output_path_injection(tmp_path):
    j=job(b'secret','text/plain','docx'); j['input_sha256']='0'*64
    with pytest.raises(ValueError): converter.convert(j)
    # There is no output path protocol: symlink/traversal never reach a file write.
    sentinel=tmp_path/'sentinel';sentinel.write_text('unchanged')
    link=tmp_path/'export.docx';link.symlink_to(sentinel)
    j=job(b'secret','text/plain','docx');j['output_path']=str(link)
    with pytest.raises(ValueError): converter.convert(j)
    assert sentinel.read_text()=='unchanged'


def test_confinement_unavailable_no_host_fallback_no_secret_error(monkeypatch,capsys):
    def unavailable(*args,**kw): raise OSError('secret filename /home/private')
    monkeypatch.setattr(sandbox.document_sandbox,'run_confined',unavailable)
    with pytest.raises(DomainError) as exc: sandbox.convert(b'secret content','text/plain',format='docx')
    assert str(exc.value)=='source_unavailable'
    assert capsys.readouterr()==('','')


@pytest.mark.parametrize('mutation',['hash','mime','stdout','stderr','filename'])
def test_strict_protocol_wrong_output_fail_closed(monkeypatch,mutation):
    raw=b'secret content'
    value=dict(version='office-converter-v1',input_sha256=sha256(raw).hexdigest(),output=base64.b64encode(b'PK\x03\x04fake').decode(),
               sha256=sha256(b'PK\x03\x04fake').hexdigest(),mime_type=converter.MIMES['docx'],omissions=[],chart=None)
    errors=b''
    if mutation=='hash': value['sha256']='0'*64
    if mutation=='mime': value['mime_type']='text/plain'
    if mutation=='filename': value['filename']='secret.docx'
    if mutation=='stderr': errors=b'secret diagnostic'
    output=json.dumps(value).encode()+(b' secret stdout' if mutation=='stdout' else b'')
    monkeypatch.setattr(sandbox.document_sandbox,'run_confined',lambda *a,**k:(0,output,errors))
    with pytest.raises(DomainError): sandbox.convert(raw,'text/plain',format='docx')


def test_chart_contract():
    from workagent.models import new_id
    with pytest.raises(ValueError):
        ExportFile(schema_version='workagent/v1',request_id=new_id(),command_id=new_id(),version_id=new_id(),sha256='0'*64,format='png')
    with pytest.raises(ValueError):
        ChartSpec(category_column='a',value_column='b',aggregation='count',title='t',x_label='a',y_label='count')


def test_exact_base_race_performs_outside_transaction():
    from workagent.office_service import prepared_command
    from types import SimpleNamespace
    from workagent.models import new_id
    from workagent.workspace_file_models import ImportOfficeDraft, FileVersionRef
    command=ImportOfficeDraft(schema_version='workagent/v1',request_id=new_id(),command_id=new_id(),expected_work_version=1,
        source=FileVersionRef(workspace_file_id=new_id(),version_id=new_id(),sha256='0'*64))
    state={'open':False,'version':1,'retained':False}
    class DB:
        @contextmanager
        def transaction(self):
            state['open']=True
            try: yield self
            finally: state['open']=False
        def execute(self,*a): return self
        def fetchone(self): return None
    class Service:
        db=DB()
        def _scope(self,*a,**kw): pass
        def _command(self,p,ws,op,args,cmd,model,prepare,mutate):
            with self.db.transaction() as c: return mutate(c,prepare(c))
    def perform(frozen):
        assert not state['open']; state['version']=2; return b'file'
    def retain(*a): state['retained']=True
    with pytest.raises(DomainError) as exc:
        prepared_command(Service(),SimpleNamespace(id='human',kind='human'),'ws','import',{},command,dict,
                         lambda c:state['version'],perform,retain)
    assert str(exc.value)=='version_conflict' and not state['retained']


def test_unknown_tool_export_is_not_downloadable():
    from types import SimpleNamespace
    from workagent.office_service import authorize_tool_download
    class DB:
        def __init__(self,data): self.data=data
        def execute(self,*a): return self
        def fetchone(self): return self.data
    manifest=SimpleNamespace(tool_attempt_id='attempt',model_dump=lambda **kw:{'export_id':'export'})
    for row in (None,{'data':{'status':'outcome_unknown'}},{'data':{'status':'cost_only'}},
                {'data':{'status':'observed','output':{'result':{'export_id':'wrong'}}}}):
        with pytest.raises(DomainError): authorize_tool_download(DB(row),'ws',manifest)
    authorize_tool_download(DB({'data':{'status':'observed','output':{'result':{'export_id':'export'}}}}),'ws',manifest)


def test_new_office_schema_opt_in_and_legacy_tool_unchanged():
    from workagent.capabilities.office_composition import SCHEMA,enabled
    from workagent.capabilities.file_adapters import ArtifactsExport
    from workagent.capabilities.office_adapters import ArtifactsOfficeExport
    assert sha256(SCHEMA.material).hexdigest()=='37829a6d55e1b241a3eec67c51a928642fcff5597e51c6152432474beea95337'
    assert not enabled({'responses':{'continuity_version':'sustained-v1','scope_tool_sha256':'wrong'}})
    request=dict(kind='artifacts_export',input_ref={'kind':'workspace_file','id':'file','version':'version','sha256':'0'*64},format='docx')
    with pytest.raises(ValueError): ArtifactsExport.model_validate(request)
    value=ArtifactsOfficeExport.model_validate({**request,'kind':'artifacts_export_office','chart_spec':None})
    assert value.format=='docx'
