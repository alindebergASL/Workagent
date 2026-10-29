"""Controller-authored B3.2b actual TCP, restricted PG and jail checks.
Synthetic documents only. Independent readers inspect downloaded bytes.
"""
import base64
from hashlib import sha256
import io
import json
import os
from pathlib import Path
import zipfile
import pytest
from test_workspace_files_parent_http import tcp,checked,create_work,create_file,command
from test_domain import context
from workagent import office_sandbox,document_sandbox
from workagent.errors import DomainError


def export(client,ws,file,fmt,**fields):
    ref=file['current']
    result=checked(client.post(f'/v1/workspaces/{ws}/files/{file["workspace_file_id"]}/exports',json=command(
        version_id=ref['version_id'],sha256=ref['sha256'],format=fmt,**fields)),201)
    response=client.get(f'/v1/workspaces/{ws}/exports/{result["export_id"]}/download')
    assert response.status_code==200,response.text
    assert response.headers['x-content-type-options']=='nosniff'
    assert response.headers['cache-control']=='no-store'
    assert response.headers['content-type'].split(';')[0]==result['blob']['mime_type']
    assert response.headers['x-content-sha256']==result['blob']['sha256']==sha256(response.content).hexdigest()
    assert len(response.content)==result['blob']['byte_length']
    assert response.headers['content-disposition']=='attachment; filename="export.'+fmt+'"'
    out=os.environ.get('OFFICE_EVIDENCE_DIR')
    if out:
        dest=Path(out);dest.mkdir(parents=True,exist_ok=True)
        (dest/(result['export_id']+'.'+fmt)).write_bytes(response.content)
        (dest/(result['export_id']+'.json')).write_text(json.dumps(result,indent=2))
    return result,response.content


def test_tcp_all_binary_outputs_independent_read_restart(tcp):
    from docx import Document
    from openpyxl import load_workbook
    from pypdf import PdfReader
    from PIL import Image
    client,ctx,start,stop,tmp=tcp;s,p,ws,_,_=ctx;cv=create_work(client,ws)
    content=dict(office_version='office-v1',kind='document',paragraphs=['Quarterly operations note','Exact quotation: “Preserve this wording.”','A measured draft, not an approved decision.'],annotations=['Human annotation: protect this exact sentence.'])
    file=create_file(client,ws,cv,json.dumps(content,ensure_ascii=False),'PRIVATE-FILENAME-SENTINEL.json','application/json')
    artifacts=[]
    for fmt in ('docx','pdf'):
        manifest,raw=export(client,ws,file,fmt);artifacts.append((manifest,raw))
        if fmt=='docx': text='\n'.join(p.text for p in Document(io.BytesIO(raw)).paragraphs)
        else:text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(raw)).pages)
        assert 'Preserve this wording.' in text and 'protect this exact sentence.' in text
        assert 'PRIVATE-FILENAME-SENTINEL' not in json.dumps(manifest) and b'PRIVATE-FILENAME-SENTINEL' not in raw
    data='region,hours\nEast,0.1\nEast,0.2\nWest,1.5\n'
    table=create_file(client,ws,cv,data,'capacity.csv','text/csv')
    manifest,raw=export(client,ws,table,'xlsx');artifacts.append((manifest,raw))
    book=load_workbook(io.BytesIO(raw),data_only=False)
    assert book['Data']['A1'].value=='region' and book['Data']['B2'].value=='0.1'
    chart=dict(category_column='region',value_column='hours',aggregation='sum',title='Computed hours by region',x_label='Region',y_label='Hours')
    manifest,raw=export(client,ws,table,'png',chart_spec=chart);artifacts.append((manifest,raw))
    image=Image.open(io.BytesIO(raw));image.verify();assert image.size==(1000,600)
    assert [(r['category'],r['value']) for r in manifest['chart']['rows']]==[('East','0.3'),('West','1.5')]
    assert manifest['chart']['input_sha256']==table['current']['sha256']
    stop();start()
    for manifest,raw in artifacts:
        again=client.get(f'/v1/workspaces/{ws}/exports/{manifest["export_id"]}/download');assert again.content==raw
    response=client.get(f'/v1/workspaces/{ws}/files/{file["workspace_file_id"]}/versions/{file["current"]["version_id"]}',params={'sha256':file['current']['sha256']})
    assert json.loads(checked(response)['text'])==content
    unauthorized=client.get(f'/v1/workspaces/{ws}/exports/{artifacts[0][0]["export_id"]}/download',headers={'Authorization':'Bearer wrong'})
    assert unauthorized.status_code==401


def test_tcp_formula_vs_string_and_cached_evidence(tcp):
    from openpyxl import Workbook,load_workbook
    client,ctx,start,stop,tmp=tcp;s,p,ws,_,_=ctx;cv=create_work(client,ws)
    wb=Workbook();sheet=wb.active;sheet.title='Values';sheet['A1']=2;sheet['A2']=3;sheet['A3']='=SUM(A1:A2)';sheet['B1']='=HYPERLINK("https://invalid.example","not executable")';sheet['B1'].data_type='s'
    buf=io.BytesIO();wb.save(buf);source=zipfile.ZipFile(io.BytesIO(buf.getvalue()));out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as dest:
        for item in source.infolist():
            raw=source.read(item.filename)
            if item.filename=='xl/worksheets/sheet1.xml':raw=raw.replace(b'<f>SUM(A1:A2)</f><v></v>',b'<f>SUM(A1:A2)</f><v>5</v>')
            dest.writestr(item,raw)
    original=out.getvalue()
    # Runtime document extraction uses its existing pinned jail; preserve original bytes.
    detail=checked(client.post(f'/v1/workspaces/{ws}/conversations/{cv["id"]}/documents',json=command(filename='formula-source.xlsx',mime_type=office_sandbox.FORMATS['xlsx'][1],content_base64=base64.b64encode(original).decode(),sha256=sha256(original).hexdigest())),201)
    file=checked(client.post(f'/v1/workspaces/{ws}/work-areas/{cv["id"]}/files',json=command(expected_work_version=cv['work_version'],source=dict(kind='document',document_id=detail['document_id'],sha256=detail['sha256']))),201)
    draft=checked(client.post(f'/v1/workspaces/{ws}/work-areas/{cv["id"]}/office-drafts',json=command(expected_work_version=cv['work_version'],source=file['current'])),201)
    ref=draft['current'];path=f'/v1/workspaces/{ws}/files/{draft["workspace_file_id"]}/versions/{ref["version_id"]}'
    representation=json.loads(checked(client.get(path,params={'sha256':ref['sha256']}))['text'])
    cells={x['address']:x for x in representation['sheets'][0]['cells']}
    assert cells['A3']['kind']=='formula' and cells['A3']['value']=='=SUM(A1:A2)' and cells['A3']['cached_value']=='5'
    assert cells['B1']['kind']=='string' and cells['B1']['value'].startswith('=HYPERLINK')
    representation['annotations'].append('Human annotation after import')
    content=json.dumps(representation)
    edited=checked(client.post(f'/v1/workspaces/{ws}/files/{draft["workspace_file_id"]}/versions',json=command(expected_work_version=cv['work_version'],base_version_id=ref['version_id'],base_sha256=ref['sha256'],content=content,sha256=sha256(content.encode()).hexdigest())),201)
    manifest,raw=export(client,ws,edited,'xlsx')
    parsed=load_workbook(io.BytesIO(raw),data_only=False)
    assert parsed['Values']['A3'].data_type=='f' and parsed['Values']['A3'].value=='=SUM(A1:A2)'
    assert parsed['Values']['B1'].data_type=='s'
    assert load_workbook(io.BytesIO(raw),data_only=True)['Values']['A3'].value is None
    assert any('cache' in o.lower() for o in manifest['omissions'])
    assert manifest['provenance']['original_sha256']==sha256(original).hexdigest()


def test_actual_jail_rejects_zip_abuse_and_formula_injection():
    from openpyxl import Workbook
    from docx import Document
    for payload in ('=WEBSERVICE("https://invalid.example")','=HYPERLINK("https://invalid.example","x")',"='[external.xlsx]Sheet1'!A1"):
        value=dict(office_version='office-v1',kind='workbook',sheets=[dict(name='Data',cells=[dict(address='A1',kind='formula',value=payload,cached_value=None)])],annotations=[])
        with pytest.raises(DomainError):office_sandbox.convert(json.dumps(value).encode(),'application/json',format='xlsx')
    doc=Document();doc.add_paragraph('valid baseline');buf=io.BytesIO();doc.save(buf)
    for filename,payload in [('word/oversize.txt',b'A'*3000000),('../traversal.txt',b'sentinel')]:
        out=io.BytesIO(buf.getvalue())
        with zipfile.ZipFile(out,'a',zipfile.ZIP_DEFLATED) as z:z.writestr(filename,payload)
        with pytest.raises(DomainError):office_sandbox.convert(out.getvalue(),office_sandbox.FORMATS['docx'][1],operation='import')


def test_actual_jail_output_symlink_cannot_escape_and_no_secrets(tmp_path,monkeypatch):
    sentinel=tmp_path/'PRIVATE-SENTINEL';sentinel.write_text('private unchanged')
    monkeypatch.setenv('OFFICE_SECRET_SENTINEL','not inherited')
    module=tmp_path/'probe.py'
    module.write_text('''import os,pathlib,json
assert 'OFFICE_SECRET_SENTINEL' not in os.environ
assert not pathlib.Path('/home/ubuntu').exists()
p=pathlib.Path('/tmp/export.docx');p.symlink_to('/home/ubuntu/PRIVATE-SENTINEL')
try:p.write_text('escape');raise AssertionError('escape succeeded')
except OSError:pass
print(json.dumps({'escape':False}))
''')
    code,output,errors=document_sandbox.run_confined(b'probe','office',module_path=module,packages=office_sandbox.package_dir(),launcher=office_sandbox.jail_argv)
    assert code==0 and json.loads(output)=={'escape':False} and not errors
    assert sentinel.read_text()=='private unchanged'
    with pytest.raises(DomainError) as exc:office_sandbox.convert(b'PRIVATE-SENTINEL','application/vnd.openxmlformats-officedocument.wordprocessingml.document',operation='import')
    assert str(exc.value)=='source_unavailable' and 'PRIVATE' not in str(exc.value)
