"""Parent-run restricted PG integration. Real converter, no fake binary success.

Requires migrated disposable DATABASE_URL/MIGRATION_DATABASE_URL and pinned
WORKAGENT_OFFICE_PACKAGES. Parent document environment is selected explicitly.
"""
import base64
from hashlib import sha256
import io
import json
import os
from pathlib import Path
import pytest
from test_domain import context,cmd,raises
from workagent.models import CreateConversation,IngestDocument
from workagent.workspace_file_models import CreateWorkspaceFile,DocumentEvidence,ImportOfficeDraft,ExportFile,EditWorkspaceFile,ProposeFilePatch
from workagent.db import Database
from workagent import document_sandbox,office_sandbox


def imported(context,monkeypatch):
    from docx import Document
    s,p,ws,refs,_=context
    cv=s.create_conversation(p,ws,cmd(CreateConversation,selected_source_refs=refs[:1]))
    doc=Document();doc.add_paragraph('Attributed quotation “unchanged”.');stream=io.BytesIO();doc.save(stream)
    raw=stream.getvalue()
    packages=Path(os.environ['WORKAGENT_DOCUMENT_PACKAGES'])
    monkeypatch.setattr(document_sandbox,'package_dir',lambda:packages)
    detail=s.ingest_document(p,ws,cv.id,cmd(IngestDocument,filename='original.docx',mime_type=office_sandbox.FORMATS['docx'][1],
        content_base64=base64.b64encode(raw).decode(),sha256=sha256(raw).hexdigest()))
    file=s.create_workspace_file(p,ws,cv.id,cmd(CreateWorkspaceFile,expected_work_version=cv.work_version,
        source=DocumentEvidence(kind='document',document_id=detail.document_id,sha256=detail.sha256)))
    command=cmd(ImportOfficeDraft,expected_work_version=cv.work_version,source=file.current)
    draft=s.import_office_draft(p,ws,cv.id,command)
    assert s.import_office_draft(p,ws,cv.id,command)==draft
    return cv,file,draft,detail,raw


def test_real_import_edit_export_revoke_original_dependency(context,monkeypatch):
    from docx import Document
    s,p,ws,refs,admin=context
    cv,original,draft,detail,raw=imported(context,monkeypatch)
    ref=draft.current
    content=json.loads(s.read_workspace_file(p,ws,draft.workspace_file_id,ref.version_id,ref.sha256,limit=8192).text)
    content['annotations'].append('Human annotation kept')
    content=json.dumps(content,ensure_ascii=False)
    draft=s.edit_workspace_file(p,ws,draft.workspace_file_id,cmd(EditWorkspaceFile,expected_work_version=cv.work_version,
        base_version_id=ref.version_id,base_sha256=ref.sha256,content=content,sha256=sha256(content.encode()).hexdigest()))
    raises('version_conflict',lambda:s.propose_file_patch(p,ws,draft.workspace_file_id,cmd(ProposeFilePatch,
        expected_work_version=cv.work_version,base_version_id=ref.version_id,base_sha256=ref.sha256,
        content=content,sha256=sha256(content.encode()).hexdigest(),reason='stale')))
    exported=s.export_workspace_file(p,ws,draft.workspace_file_id,cmd(ExportFile,version_id=draft.current.version_id,sha256=draft.current.sha256,format='docx'))
    manifest,output=s.download_export(p,ws,exported.export_id)
    assert 'Human annotation: Human annotation kept' in [p.text for p in Document(io.BytesIO(output)).paragraphs]
    assert s.download_document(p,ws,cv.id,detail.document_id)[1]==raw
    assert manifest.provenance.original_sha256==sha256(raw).hexdigest()
    with Database(admin).transaction() as c:
        c.execute('UPDATE source_access SET active=false WHERE workspace_id=%s AND principal_id=%s AND source_id=%s',(ws,p.id,refs[0].source_id))
    raises('not_found_or_not_authorized',lambda:s.download_export(p,ws,exported.export_id))
    raises('not_found_or_not_authorized',lambda:s.get_artifact(p,ws,draft.workspace_file_id))
    with s.db.transaction() as c:
        assert not c.execute('SELECT office_source_visible(%s,%s,%s) AS ok',(ws,p.id,original.current.version_id)).fetchone()['ok']
        assert not c.execute('SELECT work_artifact_visible(%s,%s,%s,%s) AS ok',(ws,p.id,cv.id,draft.workspace_file_id)).fetchone()['ok']


def test_real_conversion_race_never_retains_over_human_edit(context,monkeypatch):
    s,p,ws,_,_=context
    cv,_,draft,_,_=imported(context,monkeypatch)
    real=office_sandbox.convert
    def concurrent_edit(*args,**kwargs):
        content='{"office_version":"office-v1","kind":"document","paragraphs":["Human wins"],"annotations":[]}'
        s.edit_workspace_file(p,ws,draft.workspace_file_id,cmd(EditWorkspaceFile,expected_work_version=cv.work_version,
            base_version_id=draft.current.version_id,base_sha256=draft.current.sha256,content=content,sha256=sha256(content.encode()).hexdigest()))
        return real(*args,**kwargs)
    monkeypatch.setattr(office_sandbox,'convert',concurrent_edit)
    raises('version_conflict',lambda:s.export_workspace_file(p,ws,draft.workspace_file_id,cmd(ExportFile,
        version_id=draft.current.version_id,sha256=draft.current.sha256,format='docx')))
    with s.db.transaction() as c:
        assert c.execute('SELECT count(*) AS n FROM artifact_exports WHERE workspace_id=%s',(ws,)).fetchone()['n']==0


def test_run_office_chart_exact_receipt_replay_and_human_download(context):
    from test_workspace_files_pg import make
    from test_workspace_file_gateway_pg import install,call
    from workagent.capabilities.models import InputRef
    from workagent.capabilities.broker import CapabilityBroker
    from workagent.workspace_file_models import WorkingText
    from workagent.models import new_id
    s,p,ws,_,admin=context
    cv=s.create_conversation(p,ws,cmd(CreateConversation,selected_source_refs=[]))
    content='team,hours\na,0.1\na,0.2\nb,1.5\n'
    file=s.create_workspace_file(p,ws,cv.id,cmd(CreateWorkspaceFile,expected_work_version=cv.work_version,
        source=WorkingText(filename='dataset.csv',mime_type='text/csv',content=content,sha256=sha256(content.encode()).hexdigest())))
    ref=InputRef(kind='workspace_file',id=file.workspace_file_id,version=file.current.version_id,sha256=file.current.sha256)
    cap,grant,q=install(context,cv,[ref],['artifacts.export.office'])
    args=dict(kind='artifacts_export_office',input_ref=ref.model_dump(),format='png',chart_spec=dict(
        category_column='team',value_column='hours',aggregation='sum',title='Hours by team',x_label='Team',y_label='Hours'))
    request=call(cap,q,'artifacts.export.office',args,[ref])
    broker=CapabilityBroker(s,cap)
    receipt=broker.call(request)
    assert receipt.status=='observed',receipt.status
    assert broker.call(request)==receipt
    manifest,raw=s.download_export(p,ws,receipt.output['result']['export_id'])
    assert manifest.filename=='export.png' and manifest.tool_attempt_id==receipt.attempt_id
    assert manifest.chart.rows[0].value=='0.3'
    assert manifest.chart.input_sha256==ref.sha256 and raw.startswith(b'\x89PNG')
    # Raw SQL resource check covers exact selectors and dependency wrappers too.
    with s.db.transaction() as c:
        assert c.execute('SELECT tool_file_resources_authorized(%s,%s,%s) AS ok',
            (ws,q.run.id,__import__('workagent.service',fromlist=['encoded']).encoded(request))).fetchone()['ok']
    with Database(admin).transaction() as c:
        c.execute('UPDATE tool_grants SET active=false WHERE id=%s',(grant.id,))
    from workagent.errors import DomainError
    with pytest.raises(DomainError): broker.call(request.model_copy(update={'call_id':new_id()}))
